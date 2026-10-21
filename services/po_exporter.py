import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def is_case_uom(uom_str):
    if not uom_str:
        return False
    u = str(uom_str).strip().lower().rstrip('.')
    case_keywords = {'case', 'cases', 'box', 'boxes', 'carton', 'cartons', 'កេស', 'កាតុង'}
    return u in case_keywords

def process_po_lines_for_export(lines):
    """
    Processes PO line items for export without prices.
    Calculates case conversions, piece totals, and footer summaries.
    """
    processed = []
    has_case = False
    total_cases = 0.0
    total_pcs = 0.0

    for idx, l in enumerate(lines, 1):
        qty = float(l.get('quantity') or 0)
        uom = str(l.get('uom') or l.get('line_uom') or 'pcs.').strip()
        is_case = is_case_uom(uom)
        if is_case:
            has_case = True

        # Calculate ratio to pieces
        ratio = float(l.get('purchasing_uom_ratio_std') or 0)
        if ratio <= 1.0:
            ratio = float(l.get('sales_uom_ratio_std') or 0)
        if ratio <= 1.0:
            ratio = float(l.get('purchasing_ratio') or 0)
        if ratio <= 1.0:
            ratio = float(l.get('sales_ratio') or 0)
        if ratio <= 1.0:
            packs = float(l.get('packs_per_case') or 0)
            pcs = float(l.get('pcs_per_pack') or 0)
            if packs > 0 and pcs > 0:
                ratio = packs * pcs
        if ratio <= 0:
            ratio = 1.0

        if is_case:
            line_cases = qty
            line_pcs = qty * ratio
        else:
            line_cases = 0.0
            line_pcs = qty

        total_cases += line_cases
        total_pcs += line_pcs

        processed.append({
            'line_num': l.get('line_num') or idx,
            'product_name': l.get('product_name') or l.get('description') or '',
            'sku': l.get('product_sku') or l.get('sku') or '',
            'vendor_code': l.get('vendor_product_code') or '',
            'quantity': qty,
            'uom': uom,
            'is_case': is_case,
            'ratio': ratio,
            'line_cases': line_cases,
            'total_pcs': line_pcs
        })

    return processed, has_case, total_cases, total_pcs

def generate_po_excel(order, raw_lines, company=None):
    """
    Generates an Excel workbook for Purchase Order Packing / Receiving list (No Prices).
    Includes UOM, Total in Pcs (if cases), and bottom summary for total cases and pcs.
    """
    lines, has_case, total_cases, total_pcs = process_po_lines_for_export(raw_lines)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Order"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    font_company = Font(name="Segoe UI", size=15, bold=True, color="7A5531")
    font_sub = Font(name="Segoe UI", size=9, color="64748B")
    font_title = Font(name="Segoe UI", size=13, bold=True, color="1E293B")
    font_meta_lbl = Font(name="Segoe UI", size=9.5, bold=True, color="475569")
    font_meta_val = Font(name="Segoe UI", size=9.5, color="0F172A")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9.5, color="1E293B")
    font_data_bold = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    font_total_lbl = Font(name="Segoe UI", size=10, bold=True, color="7A5531")
    font_total_val = Font(name="Segoe UI", size=11, bold=True, color="0F172A")

    fill_header = PatternFill(start_color="7A5531", end_color="7A5531", fill_type="solid")
    fill_zebra = PatternFill(start_color="FAF7F2", end_color="FAF7F2", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_total = PatternFill(start_color="F5EFEB", end_color="F5EFEB", fill_type="solid")

    thin_side = Side(style='thin', color='E2E8F0')
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(style='medium', color='52361C'))
    border_total = Border(top=Side(style='thin', color='7A5531'), bottom=Side(style='double', color='7A5531'), left=thin_side, right=thin_side)

    align_left = Alignment(horizontal='left', vertical='center')
    align_center = Alignment(horizontal='center', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # 1. Company Header
    comp_name = (company.get('name') if company else '') or 'inFlow Premium Store'
    comp_addr = (company.get('address1') if company else '') or 'Phnom Penh, Cambodia'
    comp_phone = (company.get('phone') if company else '') or '+855 23 888 999'

    ws["A1"] = comp_name
    ws["A1"].font = font_company
    ws["A2"] = f"Address: {comp_addr}  |  Phone: {comp_phone}"
    ws["A2"].font = font_sub

    # 2. Document Title
    ws["A4"] = "PURCHASE ORDER (PACKING & RECEIVING LIST)"
    ws["A4"].font = font_title

    # 3. Order Meta Information
    order_no = order.get('order_number') or ''
    order_date = order.get('order_date') or ''
    if hasattr(order_date, 'strftime'):
        order_date_str = order_date.strftime('%b %d, %Y')
    else:
        order_date_str = str(order_date)

    vendor_name = order.get('vendor_name') or 'Queen Pack'
    vendor_phone = order.get('phone') or order.get('vendor_phone') or '-'
    vendor_addr = order.get('vendor_address') or '-'
    loc_name = order.get('location_name') or 'Main Warehouse'
    carrier = order.get('carrier') or 'Standard Carrier'

    ws["A6"] = "Order Number:"
    ws["A6"].font = font_meta_lbl
    ws["B6"] = order_no
    ws["B6"].font = font_meta_val

    ws["D6"] = "Order Date:"
    ws["D6"].font = font_meta_lbl
    ws["E6"] = order_date_str
    ws["E6"].font = font_meta_val

    ws["A7"] = "Vendor / Supplier:"
    ws["A7"].font = font_meta_lbl
    ws["B7"] = vendor_name
    ws["B7"].font = font_meta_val

    ws["D7"] = "Warehouse Destination:"
    ws["D7"].font = font_meta_lbl
    ws["E7"] = loc_name
    ws["E7"].font = font_meta_val

    ws["A8"] = "Vendor Phone:"
    ws["A8"].font = font_meta_lbl
    ws["B8"] = vendor_phone
    ws["B8"].font = font_meta_val

    ws["D8"] = "Carrier / Method:"
    ws["D8"].font = font_meta_lbl
    ws["E8"] = carrier
    ws["E8"].font = font_meta_val

    # 4. Table Header
    curr_row = 10
    if has_case:
        headers = ["#", "Product Item", "SKU", "Quantity", "UOM", "Total in Pcs"]
    else:
        headers = ["#", "Product Item", "SKU", "Quantity", "UOM"]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.border = border_header
        if h in ("Quantity", "Total in Pcs"):
            cell.alignment = align_right
        elif h in ("#", "UOM"):
            cell.alignment = align_center
        else:
            cell.alignment = align_left
    ws.row_dimensions[curr_row].height = 24

    # 5. Table Data Rows (NO PRICES, NO VENDOR CODE)
    for idx, l in enumerate(lines, 1):
        curr_row += 1
        ws.row_dimensions[curr_row].height = 20
        fill = fill_white if idx % 2 != 0 else fill_zebra

        c1 = ws.cell(row=curr_row, column=1, value=idx)
        c1.alignment = align_center

        c2 = ws.cell(row=curr_row, column=2, value=l['product_name'])
        c2.alignment = align_left

        c3 = ws.cell(row=curr_row, column=3, value=l['sku'])
        c3.alignment = align_left

        c4 = ws.cell(row=curr_row, column=4, value=l['quantity'])
        c4.alignment = align_right
        c4.number_format = '#,##0.00' if (l['quantity'] % 1 != 0) else '#,##0'

        c5 = ws.cell(row=curr_row, column=5, value=l['uom'])
        c5.alignment = align_center

        row_cells = [c1, c2, c3, c4, c5]

        if has_case:
            c6 = ws.cell(row=curr_row, column=6, value=l['total_pcs'])
            c6.alignment = align_right
            c6.font = font_data_bold
            c6.number_format = '#,##0.00' if (l['total_pcs'] % 1 != 0) else '#,##0'
            row_cells.append(c6)

        for cell in row_cells:
            cell.border = border_cell
            cell.fill = fill
            if not has_case or cell != row_cells[-1]:
                cell.font = font_data

    # 6. Bottom Totals Summary
    curr_row += 1
    ws.row_dimensions[curr_row].height = 24
    num_cols = len(headers)

    if has_case:
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        c_lbl = ws.cell(row=curr_row, column=1, value="TOTAL SUMMARY (សរុបចំនួនកេស និងចំនួនគ្រាប់):")
        c_lbl.font = font_total_lbl
        c_lbl.alignment = Alignment(horizontal='right', vertical='center')

        # Total Cases
        c_cases = ws.cell(row=curr_row, column=4, value=f"{total_cases:,.0f} cases" if total_cases % 1 == 0 else f"{total_cases:,.2f} cases")
        c_cases.font = font_total_val
        c_cases.alignment = align_right

        # Blank UOM column cell
        c_uom_space = ws.cell(row=curr_row, column=5, value="")

        # Total Pcs
        c_pcs = ws.cell(row=curr_row, column=6, value=f"{total_pcs:,.0f} pcs" if total_pcs % 1 == 0 else f"{total_pcs:,.2f} pcs")
        c_pcs.font = font_total_val
        c_pcs.alignment = align_right

        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.fill = fill_total
            cell.border = border_total
    else:
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        c_lbl = ws.cell(row=curr_row, column=1, value="TOTAL PIECES (ចំនួនសរុប):")
        c_lbl.font = font_total_lbl
        c_lbl.alignment = Alignment(horizontal='right', vertical='center')

        c_pcs = ws.cell(row=curr_row, column=4, value=f"{total_pcs:,.0f} pcs" if total_pcs % 1 == 0 else f"{total_pcs:,.2f} pcs")
        c_pcs.font = font_total_val
        c_pcs.alignment = align_right

        for col_idx in range(1, num_cols + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.fill = fill_total
            cell.border = border_total

    # Column Widths (Wider Product Item column)
    col_widths = {
        1: 6,
        2: 60,
        3: 20,
        4: 15,
        5: 12,
        6: 18
    }
    for c_idx, w in col_widths.items():
        if c_idx <= num_cols:
            ws.column_dimensions[get_column_letter(c_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
