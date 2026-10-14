import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    """Canvas that adds page numbers and running footer to ReportLab PDFs."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Header rule & text
        self.drawString(24, 580, "inFlow Premium • Monthly Expense & Cash Flow Report")
        self.drawRightString(768, 580, f"Page {self._pageNumber} of {page_count}")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(24, 574, 768, 574)
        
        # Footer rule & text
        self.line(24, 28, 768, 28)
        self.drawString(24, 18, f"Confidential & Proprietary • inFlow Premium ERP • Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        self.drawRightString(768, 18, "www.inflowinventory.com")
        self.restoreState()


def generate_monthly_excel(data):
    """
    Generates a beautifully formatted Excel (.xlsx) workbook for the monthly report.
    Matches the user's Google Sheet structure with KPI cards, expenses, accounts, and media receipts.
    """
    wb = openpyxl.Workbook()
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    double_bottom_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='double', color='0F766E')
    )

    month_display = data.get('month_display', '')
    month_key = data.get('month_key', '')

    # -------------------------------------------------------------------------
    # SHEET 1: Monthly Summary & Daily Expenses
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Expenses & Cash Reconciliation"
    ws1.views.sheetView[0].showGridLines = True

    # 1. Main Title Banner
    ws1.merge_cells("A1:I1")
    title_cell = ws1["A1"]
    title_cell.value = f"inFlow Premium - Monthly Financial & Expense Report ({month_display})"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    # Subtitle
    ws1.merge_cells("A2:I2")
    sub_cell = ws1["A2"]
    sub_cell.value = f"Report Period: {month_display} ({month_key}) | Exported: {datetime.now().strftime('%d-%b-%Y %H:%M')} | Currency: USD ($)"
    sub_cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # 2. Executive KPI Summary Row (Matching Google Sheet Top Row)
    is_branch = data.get('is_branch', False)
    if is_branch:
        kpi_headers = [
            ("ចំណាយសរុប (Expenses)", "A4", "FEE2E2", "DC2626"),
            ("លុយចាស់ (Opening)", "B4", "DCFCE7", "166534"),
            ("លុយលក់ (Customer Receipts)", "C4", "CCFBF1", "0F766E"),
            ("លុយសរុប (Total Inflow)", "D4", "DBEAFE", "1D4ED8"),
            ("លុយសល់ (Money in Hand)", "E4", "FEF9C3", "B45309"),
            ("ចំណូល - ចំណាយ (Net Margin)", "F4", "DCFCE7", "15803D")
        ]
        kpi_values = [
            data.get('total_expenses', 0.0),
            data.get('opening_balance', 0.0),
            data.get('sales_received', 0.0),
            data.get('total_money_in', 0.0),
            data.get('money_in_hand', 0.0),
            data.get('net_operating_income', 0.0)
        ]
        cols = ['A', 'B', 'C', 'D', 'E', 'F']
    else:
        kpi_headers = [
            ("ចំណាយសរុប (Expenses)", "A4", "FEE2E2", "DC2626"),
            ("លុយចាស់ (Opening)", "B4", "DCFCE7", "166534"),
            ("Clear STM", "C4", "E0F2FE", "0369A1"),
            ("Media (ប្រាក់ពី Media)", "D4", "DCFCE7", "15803D"),
            ("លុយលក់ (Customer Receipts)", "E4", "CCFBF1", "0F766E"),
            ("លុយសរុប (Total Inflow)", "F4", "DBEAFE", "1D4ED8"),
            ("លុយសល់ (Money in Hand)", "G4", "FEF9C3", "B45309"),
            ("ចំណូល - ចំណាយ (Net Margin)", "H4", "DCFCE7", "15803D")
        ]
        kpi_values = [
            data.get('total_expenses', 0.0),
            data.get('opening_balance', 0.0),
            data.get('clear_stm', 0.0),
            data.get('media_total', 0.0),
            data.get('sales_received', 0.0),
            data.get('total_money_in', 0.0),
            data.get('money_in_hand', 0.0),
            data.get('net_operating_income', 0.0)
        ]
        cols = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']

    ws1.row_dimensions[4].height = 22
    ws1.row_dimensions[5].height = 28

    for idx, (label, cell_ref, fill_color, text_color) in enumerate(kpi_headers):
        col = cols[idx]
        h_cell = ws1[f"{col}4"]
        h_cell.value = label
        h_cell.font = Font(name="Segoe UI", size=8.5, bold=True, color=text_color)
        h_cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        h_cell.alignment = Alignment(horizontal="center", vertical="center")
        h_cell.border = thin_border

        v_cell = ws1[f"{col}5"]
        v_cell.value = kpi_values[idx]
        v_cell.number_format = '$#,##0.00'
        v_cell.font = Font(name="Segoe UI", size=12, bold=True, color=text_color)
        v_cell.alignment = Alignment(horizontal="center", vertical="center")
        v_cell.border = thin_border
        v_cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

    # 3. Daily Expenses Table Section
    ws1["A7"] = "DAILY EXPENSE BREAKDOWN (តារាងចំណាយប្រចាំថ្ងៃ)"
    ws1["A7"].font = Font(name="Segoe UI", size=11, bold=True, color="0F766E")
    ws1.row_dimensions[7].height = 24

    expense_headers = ["No", "Item Description (មុខទំនិញ)", "Unit Price ($)", "Qty", "Total Price ($)", "Date", "Paid By", "Category", "Notes"]
    ws1.row_dimensions[8].height = 24
    for c_idx, h_text in enumerate(expense_headers, 1):
        c = ws1.cell(row=8, column=c_idx, value=h_text)
        c.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        c.alignment = Alignment(horizontal="center" if c_idx in [1, 4, 6] else ("right" if c_idx in [3, 5] else "left"), vertical="center")
        c.border = thin_border

    curr_row = 9
    expenses = data.get('expenses', [])
    for idx, exp in enumerate(expenses, 1):
        ws1.row_dimensions[curr_row].height = 20
        row_fill = PatternFill(start_color="F0FDFA" if idx % 2 == 0 else "FFFFFF", end_color="F0FDFA" if idx % 2 == 0 else "FFFFFF", fill_type="solid")
        
        c_no = ws1.cell(row=curr_row, column=1, value=idx)
        c_no.alignment = Alignment(horizontal="center", vertical="center")
        
        c_item = ws1.cell(row=curr_row, column=2, value=exp.get('item', ''))
        c_item.alignment = Alignment(horizontal="left", vertical="center")
        c_item.font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
        
        c_price = ws1.cell(row=curr_row, column=3, value=float(exp.get('unit_price', 0) or 0))
        c_price.number_format = '$#,##0.00'
        c_price.alignment = Alignment(horizontal="right", vertical="center")
        
        c_qty = ws1.cell(row=curr_row, column=4, value=float(exp.get('quantity', 1) or 1))
        c_qty.number_format = '#,##0.##'
        c_qty.alignment = Alignment(horizontal="center", vertical="center")
        
        c_tot = ws1.cell(row=curr_row, column=5, value=float(exp.get('total_price', 0) or 0))
        c_tot.number_format = '$#,##0.00'
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        c_tot.font = Font(name="Segoe UI", size=9.5, bold=True, color="B91C1C")
        
        exp_date = exp.get('expense_date')
        date_str = exp_date.strftime('%d-%b-%Y') if hasattr(exp_date, 'strftime') else str(exp_date or '')
        c_date = ws1.cell(row=curr_row, column=6, value=date_str)
        c_date.alignment = Alignment(horizontal="center", vertical="center")
        
        c_payer = ws1.cell(row=curr_row, column=7, value=exp.get('paid_by', ''))
        c_payer.alignment = Alignment(horizontal="left", vertical="center")
        
        c_cat = ws1.cell(row=curr_row, column=8, value=exp.get('category', ''))
        c_cat.alignment = Alignment(horizontal="left", vertical="center")
        
        c_notes = ws1.cell(row=curr_row, column=9, value=exp.get('notes', ''))
        c_notes.alignment = Alignment(horizontal="left", vertical="center")

        for col_i in range(1, 10):
            cell = ws1.cell(row=curr_row, column=col_i)
            cell.border = thin_border
            cell.fill = row_fill
            if col_i not in [2, 5]:
                cell.font = Font(name="Segoe UI", size=9)

        curr_row += 1

    # Total Expenses Row
    ws1.row_dimensions[curr_row].height = 24
    ws1.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
    tot_lbl = ws1.cell(row=curr_row, column=1, value=f"សរុបការចំណាយ (Total {len(expenses)} Items):")
    tot_lbl.font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
    tot_lbl.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    tot_val = ws1.cell(row=curr_row, column=5, value=data.get('total_expenses', 0.0))
    tot_val.number_format = '$#,##0.00'
    tot_val.font = Font(name="Segoe UI", size=11, bold=True, color="B91C1C")
    tot_val.alignment = Alignment(horizontal="right", vertical="center")
    tot_val.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    for col_i in range(1, 10):
        c = ws1.cell(row=curr_row, column=col_i)
        c.border = double_bottom_border
        if col_i > 5:
            c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    # 4. Bank & Cash Accounts Reconciliation Table
    curr_row += 3
    ws1.cell(row=curr_row, column=1, value="CASH & BANK ACCOUNTS RECONCILIATION (ផ្ទៀងផ្ទាត់សាច់ប្រាក់ក្នុងដៃ)").font = Font(name="Segoe UI", size=11, bold=True, color="0F766E")
    ws1.row_dimensions[curr_row].height = 22
    curr_row += 1

    ws1.cell(row=curr_row, column=1, value="Account / Bank Name").font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    ws1.cell(row=curr_row, column=1).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws1.cell(row=curr_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws1.cell(row=curr_row, column=1).border = thin_border

    ws1.cell(row=curr_row, column=2, value="Amount ($)").font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    ws1.cell(row=curr_row, column=2).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws1.cell(row=curr_row, column=2).alignment = Alignment(horizontal="right", vertical="center")
    ws1.cell(row=curr_row, column=2).border = thin_border

    ws1.cell(row=curr_row, column=3, value="Notes / Expression").font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
    ws1.cell(row=curr_row, column=3).fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    ws1.cell(row=curr_row, column=3).alignment = Alignment(horizontal="left", vertical="center")
    ws1.cell(row=curr_row, column=3).border = thin_border
    ws1.row_dimensions[curr_row].height = 22

    curr_row += 1
    accounts = data.get('accounts', [])
    for acc in accounts:
        ws1.row_dimensions[curr_row].height = 20
        c_n = ws1.cell(row=curr_row, column=1, value=acc.get('name', ''))
        c_n.font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
        c_n.border = thin_border

        c_a = ws1.cell(row=curr_row, column=2, value=float(acc.get('amount', 0) or 0))
        c_a.number_format = '$#,##0.00'
        c_a.font = Font(name="Segoe UI", size=9.5, bold=True, color="0F766E")
        c_a.alignment = Alignment(horizontal="right", vertical="center")
        c_a.border = thin_border

        c_e = ws1.cell(row=curr_row, column=3, value=f"Formula: {acc.get('expr')}" if acc.get('expr') and acc.get('expr') != f"{float(acc.get('amount', 0)):.2f}" else "")
        c_e.font = Font(name="Segoe UI", size=8.5, italic=True, color="64748B")
        c_e.border = thin_border
        curr_row += 1

    # Reconciliation Totals & Difference
    ws1.row_dimensions[curr_row].height = 22
    ws1.cell(row=curr_row, column=1, value="ទឹកប្រាក់សរុប (Total Counted Cash):").font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
    c_tot_count = ws1.cell(row=curr_row, column=2, value=data.get('total_counted_cash', 0.0))
    c_tot_count.number_format = '$#,##0.00'
    c_tot_count.font = Font(name="Segoe UI", size=10, bold=True, color="0F766E")
    c_tot_count.alignment = Alignment(horizontal="right", vertical="center")
    ws1.cell(row=curr_row, column=1).border = thin_border
    c_tot_count.border = thin_border
    curr_row += 1

    ws1.row_dimensions[curr_row].height = 22
    ws1.cell(row=curr_row, column=1, value="លុយសល់តាមបញ្ជី (Expected Money in Hand):").font = Font(name="Segoe UI", size=9, bold=True, color="64748B")
    c_exp_hand = ws1.cell(row=curr_row, column=2, value=data.get('money_in_hand', 0.0))
    c_exp_hand.number_format = '$#,##0.00'
    c_exp_hand.font = Font(name="Segoe UI", size=9.5, bold=True, color="B45309")
    c_exp_hand.alignment = Alignment(horizontal="right", vertical="center")
    ws1.cell(row=curr_row, column=1).border = thin_border
    c_exp_hand.border = thin_border
    curr_row += 1

    # Difference Row (Highlighted Red if negative, Green if positive/balanced)
    diff = data.get('cash_variance', 0.0)
    diff_fill = "FEE2E2" if diff < -0.005 else "DCFCE7"
    diff_text_color = "DC2626" if diff < -0.005 else "16A34A"
    diff_status = "⚠️ Cash Shortage (ខ្វះ)" if diff < -0.005 else ("🟢 Cash Surplus (លើស)" if diff > 0.005 else "✓ 100% Balanced (ត្រូវគ្នា)")

    ws1.row_dimensions[curr_row].height = 24
    lbl_diff = ws1.cell(row=curr_row, column=1, value=f"ផ្សេងៗគ្នា លុយ (Difference): {diff_status}")
    lbl_diff.font = Font(name="Segoe UI", size=9.5, bold=True, color=diff_text_color)
    lbl_diff.fill = PatternFill(start_color=diff_fill, end_color=diff_fill, fill_type="solid")
    lbl_diff.border = double_bottom_border

    val_diff = ws1.cell(row=curr_row, column=2, value=diff)
    val_diff.number_format = '$#,##0.00'
    val_diff.font = Font(name="Segoe UI", size=11, bold=True, color=diff_text_color)
    val_diff.alignment = Alignment(horizontal="right", vertical="center")
    val_diff.fill = PatternFill(start_color=diff_fill, end_color=diff_fill, fill_type="solid")
    val_diff.border = double_bottom_border

    # Auto-fit columns in Sheet 1
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip merged banner row in width calculation
            if cell.row in [1, 2]:
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------------------
    # SHEET 2: Media Receipts (ប្រាក់ទទួលបានពី Media - HQ Only)
    # -------------------------------------------------------------------------
    if not is_branch:
        ws2 = wb.create_sheet(title="Media Receipts")
        ws2.views.sheetView[0].showGridLines = True

        # Title
        ws2.merge_cells("A1:E1")
        m_title = ws2["A1"]
        m_title.value = f"inFlow Premium - Money Received from Media ({month_display})"
        m_title.font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
        m_title.fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
        m_title.alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 32

        # Headers
        media_headers = ["No", "Date (កាលបរិច្ឆេទ)", "Customer / Client Name (ឈ្មោះអតិថិជន)", "Money Received ($)", "Notes / Remark"]
        ws2.row_dimensions[3].height = 24
        for c_idx, h_text in enumerate(media_headers, 1):
            c = ws2.cell(row=3, column=c_idx, value=h_text)
            c.font = Font(name="Segoe UI", size=9.5, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="166534", end_color="166534", fill_type="solid")
            c.alignment = Alignment(horizontal="center" if c_idx in [1, 2] else ("right" if c_idx == 4 else "left"), vertical="center")
            c.border = thin_border

        media_receipts = data.get('media_receipts', [])
        m_row = 4
        for idx, m in enumerate(media_receipts, 1):
            ws2.row_dimensions[m_row].height = 20
            row_fill = PatternFill(start_color="F0FDF4" if idx % 2 == 0 else "FFFFFF", end_color="F0FDF4" if idx % 2 == 0 else "FFFFFF", fill_type="solid")

            c_no = ws2.cell(row=m_row, column=1, value=idx)
            c_no.alignment = Alignment(horizontal="center", vertical="center")

            m_date = m.get('receipt_date')
            date_str = m_date.strftime('%d-%b-%Y') if hasattr(m_date, 'strftime') else str(m_date or '')
            c_d = ws2.cell(row=m_row, column=2, value=date_str)
            c_d.alignment = Alignment(horizontal="center", vertical="center")

            c_cust = ws2.cell(row=m_row, column=3, value=m.get('customer_name', ''))
            c_cust.font = Font(name="Segoe UI", size=9.5, bold=True, color="0F172A")
            c_cust.alignment = Alignment(horizontal="left", vertical="center")

            c_amt = ws2.cell(row=m_row, column=4, value=float(m.get('amount', 0) or 0))
            c_amt.number_format = '$#,##0.00'
            c_amt.font = Font(name="Segoe UI", size=9.5, bold=True, color="16A34A")
            c_amt.alignment = Alignment(horizontal="right", vertical="center")

            c_note = ws2.cell(row=m_row, column=5, value=m.get('notes', ''))
            c_note.alignment = Alignment(horizontal="left", vertical="center")

            for col_i in range(1, 6):
                cell = ws2.cell(row=m_row, column=col_i)
                cell.border = thin_border
                cell.fill = row_fill
                if col_i not in [3, 4]:
                    cell.font = Font(name="Segoe UI", size=9)

            m_row += 1

        # Media Total Row
        ws2.row_dimensions[m_row].height = 24
        ws2.merge_cells(start_row=m_row, start_column=1, end_row=m_row, end_column=3)
        m_tot_lbl = ws2.cell(row=m_row, column=1, value=f"សរុបប្រាក់ទទួលបានពី Media (Total {len(media_receipts)} Receipts):")
        m_tot_lbl.font = Font(name="Segoe UI", size=10, bold=True, color="166534")
        m_tot_lbl.alignment = Alignment(horizontal="right", vertical="center")
        m_tot_lbl.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        m_tot_val = ws2.cell(row=m_row, column=4, value=data.get('media_total', 0.0))
        m_tot_val.number_format = '$#,##0.00'
        m_tot_val.font = Font(name="Segoe UI", size=11, bold=True, color="166534")
        m_tot_val.alignment = Alignment(horizontal="right", vertical="center")
        m_tot_val.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        for col_i in range(1, 6):
            c = ws2.cell(row=m_row, column=col_i)
            c.border = double_bottom_border
            if col_i == 5:
                c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

        for col in ws2.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == 1:
                    continue
                v_str = str(cell.value or '')
                if len(v_str) > max_len:
                    max_len = len(v_str)
            ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generate_monthly_pdf(data):
    """
    Generates a clean, professional landscape PDF report using ReportLab.
    Includes executive financial KPIs, daily expenses table, accounts count reconciliation, and media receipts.
    """
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(letter),
        leftMargin=24,
        rightMargin=24,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#0F766E')
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=8
    )

    section_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        textColor=colors.HexColor('#0F766E'),
        spaceBefore=10,
        spaceAfter=5
    )

    cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5
    )

    cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5
    )

    cell_money = ParagraphStyle(
        'TableCellMoney',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        alignment=2 # Right aligned
    )

    month_display = data.get('month_display', '')
    month_key = data.get('month_key', '')

    elements = []

    # 1. Header
    elements.append(Paragraph(f"inFlow Premium • Monthly Expense & Cash Flow Report ({month_display})", title_style))
    elements.append(Paragraph(f"Period: {month_display} ({month_key}) | Currency: USD ($) | Generated on: {datetime.now().strftime('%d-%b-%Y %H:%M')}", sub_style))

    # 2. Executive Financial KPI Summary Cards Table
    is_branch = data.get('is_branch', False)
    if is_branch:
        kpi_data = [
            [
                Paragraph("<b>TOTAL EXPENSES</b><br/><font size=10 color='#DC2626'><b>${:,.2f}</b></font>".format(data.get('total_expenses', 0.0)), cell_style),
                Paragraph("<b>OPENING BALANCE</b><br/><font size=10 color='#166534'><b>${:,.2f}</b></font>".format(data.get('opening_balance', 0.0)), cell_style),
                Paragraph("<b>SALES RECEIVED</b><br/><font size=10 color='#0F766E'><b>${:,.2f}</b></font>".format(data.get('sales_received', 0.0)), cell_style),
                Paragraph("<b>TOTAL INFLOW</b><br/><font size=10 color='#1D4ED8'><b>${:,.2f}</b></font>".format(data.get('total_money_in', 0.0)), cell_style),
                Paragraph("<b>MONEY IN HAND</b><br/><font size=10 color='#B45309'><b>${:,.2f}</b></font>".format(data.get('money_in_hand', 0.0)), cell_style),
                Paragraph("<b>NET MARGIN</b><br/><font size=10 color='#15803D'><b>${:,.2f}</b></font>".format(data.get('net_operating_income', 0.0)), cell_style)
            ]
        ]
        t_kpi = Table(kpi_data, colWidths=[124] * 6)
        t_kpi.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#FEE2E2')),
            ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#DCFCE7')),
            ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#CCFBF1')),
            ('BACKGROUND', (3, 0), (3, 0), colors.HexColor('#DBEAFE')),
            ('BACKGROUND', (4, 0), (4, 0), colors.HexColor('#FEF9C3')),
            ('BACKGROUND', (5, 0), (5, 0), colors.HexColor('#DCFCE7')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
    else:
        kpi_data = [
            [
                Paragraph("<b>TOTAL EXPENSES</b><br/><font size=10 color='#DC2626'><b>${:,.2f}</b></font>".format(data.get('total_expenses', 0.0)), cell_style),
                Paragraph("<b>OPENING BALANCE</b><br/><font size=10 color='#166534'><b>${:,.2f}</b></font>".format(data.get('opening_balance', 0.0)), cell_style),
                Paragraph("<b>CLEAR STM</b><br/><font size=10 color='#0369A1'><b>${:,.2f}</b></font>".format(data.get('clear_stm', 0.0)), cell_style),
                Paragraph("<b>MEDIA RECEIPTS</b><br/><font size=10 color='#15803D'><b>${:,.2f}</b></font>".format(data.get('media_total', 0.0)), cell_style),
                Paragraph("<b>SALES RECEIVED</b><br/><font size=10 color='#0F766E'><b>${:,.2f}</b></font>".format(data.get('sales_received', 0.0)), cell_style),
                Paragraph("<b>TOTAL INFLOW</b><br/><font size=10 color='#1D4ED8'><b>${:,.2f}</b></font>".format(data.get('total_money_in', 0.0)), cell_style),
                Paragraph("<b>MONEY IN HAND</b><br/><font size=10 color='#B45309'><b>${:,.2f}</b></font>".format(data.get('money_in_hand', 0.0)), cell_style),
                Paragraph("<b>NET MARGIN</b><br/><font size=10 color='#15803D'><b>${:,.2f}</b></font>".format(data.get('net_operating_income', 0.0)), cell_style)
            ]
        ]
        t_kpi = Table(kpi_data, colWidths=[93] * 8)
        t_kpi.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#FEE2E2')),
            ('BACKGROUND', (1, 0), (1, 0), colors.HexColor('#DCFCE7')),
            ('BACKGROUND', (2, 0), (2, 0), colors.HexColor('#E0F2FE')),
            ('BACKGROUND', (3, 0), (3, 0), colors.HexColor('#DCFCE7')),
            ('BACKGROUND', (4, 0), (4, 0), colors.HexColor('#CCFBF1')),
            ('BACKGROUND', (5, 0), (5, 0), colors.HexColor('#DBEAFE')),
            ('BACKGROUND', (6, 0), (6, 0), colors.HexColor('#FEF9C3')),
            ('BACKGROUND', (7, 0), (7, 0), colors.HexColor('#DCFCE7')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ]))
    elements.append(t_kpi)
    elements.append(Spacer(1, 8))

    # 3. Daily Expenses Breakdown
    elements.append(Paragraph(f"Daily Expense Breakdown (តារាងចំណាយប្រចាំថ្ងៃ) - Total {len(data.get('expenses', []))} items", section_style))

    exp_table_data = [
        ["No", "Item Description", "Unit Price", "Qty", "Total Price", "Date", "Paid By", "Category", "Notes"]
    ]

    for idx, exp in enumerate(data.get('expenses', []), 1):
        exp_date = exp.get('expense_date')
        date_str = exp_date.strftime('%d-%b') if hasattr(exp_date, 'strftime') else str(exp_date or '')
        exp_table_data.append([
            str(idx),
            Paragraph(exp.get('item', ''), cell_bold),
            f"${float(exp.get('unit_price', 0) or 0):,.2f}",
            f"{float(exp.get('quantity', 1) or 1):g}",
            f"${float(exp.get('total_price', 0) or 0):,.2f}",
            date_str,
            exp.get('paid_by', ''),
            exp.get('category', ''),
            Paragraph(exp.get('notes', '') or '-', cell_style)
        ])

    # Total Row
    exp_table_data.append([
        "",
        Paragraph("<b>Total Expenses:</b>", cell_style),
        "",
        "",
        f"${data.get('total_expenses', 0.0):,.2f}",
        "",
        "",
        "",
        ""
    ])

    col_widths = [24, 180, 60, 36, 70, 50, 70, 74, 180]
    t_exp = Table(exp_table_data, colWidths=col_widths, repeatRows=1)
    t_exp.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F766E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 4),
        ('TOPPADDING', (0, 0), (-1, 0), 4),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (2, 0), (2, -1), 'RIGHT'),
        ('ALIGN', (3, 0), (3, -1), 'CENTER'),
        ('ALIGN', (4, 0), (4, -1), 'RIGHT'),
        ('ALIGN', (5, 0), (5, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor('#E2E8F0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        # Total row style
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEE2E2')),
        ('TEXTCOLOR', (4, -1), (4, -1), colors.HexColor('#B91C1C')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('LINEABOVE', (0, -1), (-1, -1), 1, colors.HexColor('#DC2626')),
    ]))
    elements.append(t_exp)
    elements.append(Spacer(1, 12))

    # 4. Cash Reconciliation and Media Receipts Side-by-Side or Stacked
    recon_media_block = []

    # Accounts Reconciliation Table
    recon_rows = [
        ["Account / Bank Name", "Amount ($)"]
    ]
    for acc in data.get('accounts', []):
        recon_rows.append([
            acc.get('name', ''),
            f"${float(acc.get('amount', 0) or 0):,.2f}"
        ])
    recon_rows.append(["Total Counted Cash:", f"${data.get('total_counted_cash', 0.0):,.2f}"])
    recon_rows.append(["Expected Money in Hand:", f"${data.get('money_in_hand', 0.0):,.2f}"])
    diff = data.get('cash_variance', 0.0)
    diff_sign = "-" if diff < 0 else ("+" if diff > 0 else "")
    diff_lbl = "Difference (Shortage):" if diff < -0.005 else ("Difference (Surplus):" if diff > 0.005 else "Difference (Balanced):")
    recon_rows.append([diff_lbl, f"{diff_sign}${abs(diff):,.2f}"])

    t_recon = Table(recon_rows, colWidths=[200, 100])
    t_recon.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F766E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0, -3), (-1, -3), colors.HexColor('#F0FDFA')),
        ('BACKGROUND', (0, -2), (-1, -2), colors.HexColor('#FEF9C3')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEE2E2') if diff < -0.005 else colors.HexColor('#DCFCE7')),
        ('TEXTCOLOR', (0, -1), (-1, -1), colors.HexColor('#DC2626') if diff < -0.005 else colors.HexColor('#16A34A')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
    ]))

    if is_branch:
        elements.append(KeepTogether([
            Paragraph("<b>Cash & Bank Reconciliation</b>", section_style),
            Spacer(1, 4),
            t_recon
        ]))
    else:
        # Media Receipts Table
        media_rows = [
            ["No", "Date", "Customer Name", "Money Received ($)", "Notes"]
        ]
        for idx, m in enumerate(data.get('media_receipts', []), 1):
            m_date = m.get('receipt_date')
            date_str = m_date.strftime('%d-%b') if hasattr(m_date, 'strftime') else str(m_date or '')
            media_rows.append([
                str(idx),
                date_str,
                m.get('customer_name', ''),
                f"${float(m.get('amount', 0) or 0):,.2f}",
                m.get('notes', '') or '-'
            ])
        media_rows.append([
            "",
            Paragraph("<b>Total Media:</b>", cell_style),
            "",
            f"${data.get('media_total', 0.0):,.2f}",
            ""
        ])

        t_media = Table(media_rows, colWidths=[24, 60, 130, 80, 120])
        t_media.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#166534')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 7.5),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('ALIGN', (1, 0), (1, -1), 'CENTER'),
            ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor('#E2E8F0')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#DCFCE7')),
            ('TEXTCOLOR', (3, -1), (3, -1), colors.HexColor('#166534')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ]))

        elements.append(KeepTogether([
            Table([
                [Paragraph("<b>Cash & Bank Reconciliation</b>", section_style), Paragraph("<b>Media Receipts (ប្រាក់ទទួលបានពី Media)</b>", section_style)],
                [t_recon, t_media]
            ], colWidths=[330, 414], style=[
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('LEFTPADDING', (0, 0), (-1, -1), 0),
                ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ])
        ]))

    doc.build(elements, canvasmaker=NumberedCanvas)
    buf.seek(0)
    return buf
